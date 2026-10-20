#!/usr/bin/env python3
"""
Stage 4: collapse the seven rows on the Districts sheet that are not judicial
districts, then rebuild every total that depended on them.

Three of them are a court establishment sitting inside a larger district. Their
judges are already inside the parent's count -- the parent's figure was scraped
from the whole site, this establishment included -- so only the count
columns are carried across, then the row goes.

    Berasia          -> Bhopal            (Madhya Pradesh)
    Bheemunipatnam   -> Visakapatnam      (Andhra Pradesh)
    Parchur          -> Prakasham         (Andhra Pradesh)

The other four are district labels filed under the wrong state. Their judges
are already counted in the correct state's row, so the rows are removed
outright and the handful of accounts on them goes with them.

    Karnataka / Patna       Patna is in Bihar
    Delhi / Krishna         Krishna is in Andhra Pradesh
    Haryana / Chandigarh    Chandigarh has its own block
    Haryana / Ludhiana      Ludhiana has its own Punjab row

Rebuilt afterwards, because openpyxl does not move formulas when rows are
deleted:
  - the L/M/N ratio formulas on every district row
  - the C..K SUM ranges and L/M/N ratios on every "Total <State>" row
  - the States sheet's "Districts with accounts" and its account columns,
    both recomputed from the Districts sheet rather than adjusted by hand

    python restructure.py --in districts_filled.xlsx --out districts_final.xlsx

The MERGE, DROP and NOTES tables below are specific to one workbook -- edit
them for yours. Everything after them is generic.
"""

import argparse
import shutil

MERGE = {
    ("Madhya Pradesh", "Berasia"): "Bhopal",
    ("Andhra Pradesh", "Bheemunipatnam"): "Visakapatnam",
    ("Andhra Pradesh", "Parchur"): "Prakasham",
}
DROP = {
    ("Karnataka", "Patna"): "Patna is a Bihar district; its judges are counted "
                            "in the Bihar / Patna row",
    ("Delhi", "Krishna"): "Krishna is an Andhra Pradesh district; its judges "
                          "are counted in the Andhra Pradesh / Krishna row",
    ("Haryana", "Chandigarh"): "Chandigarh has its own block in this sheet",
    ("Haryana", "Ludhiana"): "Ludhiana is a Punjab district; its judges are "
                             "counted in the Punjab / Ludhiana row",
}
# Columns E..K hold the per-district count columns that get merged
COUNT_COLS = range(5, 12)
NOTES = {
    ("Sikkim", "Pakyong"): (
        "Not a judicial district. Pakyong was carved out of East Sikkim in "
        "2021 but has no court of its own: it falls under the Gangtok "
        "judicial district (District & Sessions Court, Gangtok). All four "
        "Gangtok establishments sit at Gangtok, so the officers serving "
        "Pakyong are counted in the Gangtok row."),
    ("Sikkim", "Soreng"): (
        "Not a judicial district. Soreng was carved out of West Sikkim in "
        "2021 but has no court of its own: it falls under the Gyalshing "
        "judicial district (District & Sessions Court, Gyalshing). All four "
        "Gyalshing establishments sit at Gyalshing, so the officers serving "
        "Soreng are counted in the Gyalshing row."),
}


def district_rows(ws):
    """(row, state, district) for the data rows, skipping the Total rows."""
    out = []
    for r in range(2, ws.max_row + 1):
        s, d = ws.cell(r, 1).value, ws.cell(r, 2).value
        if not d or str(d).startswith("Total"):
            continue
        if not s:
            continue
        out.append((r, str(s), str(d)))
    return out


def state_blocks(ws):
    """[(state, first_row, last_row, total_row)] in sheet order."""
    blocks, cur, start = [], None, None
    for r in range(2, ws.max_row + 1):
        s, d = ws.cell(r, 1).value, ws.cell(r, 2).value
        if d and str(d).startswith("Total"):
            if cur is not None:
                blocks.append((cur, start, r - 1, r))
            cur, start = None, None
            continue
        if s and d and cur is None:
            cur, start = str(s), r
    return blocks


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="src", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    from openpyxl import load_workbook
    from openpyxl.comments import Comment

    shutil.copyfile(args.src, args.out)
    wb = load_workbook(args.out, data_only=False)
    ws = wb["Districts"]

    # ---- 1. carry the count columns of each establishment into its parent
    index = {(s, d): r for r, s, d in district_rows(ws)}
    merged = []
    for (state, child), parent in MERGE.items():
        cr, pr = index.get((state, child)), index.get((state, parent))
        if cr is None or pr is None:
            raise SystemExit(f"cannot find {state}/{child} or its parent {parent}")
        moved = {}
        for c in COUNT_COLS:
            cv, pv = ws.cell(cr, c).value, ws.cell(pr, c).value
            if isinstance(cv, (int, float)) and cv:
                ws.cell(pr, c).value = (pv if isinstance(pv, (int, float)) else 0) + cv
                moved[ws.cell(1, c).value] = cv
        merged.append((state, child, parent, moved,
                       ws.cell(pr, 3).value, ws.cell(pr, 4).value))

    # ---- 2. drop both sets of rows, bottom-up so indices stay valid
    victims = []
    for key in list(MERGE) + list(DROP):
        r = index.get(key)
        if r:
            victims.append((r, key))
    dropped_counts = {}
    for r, key in sorted(victims, reverse=True):
        if key in DROP:
            dropped_counts[key] = [ws.cell(r, c).value for c in COUNT_COLS]
        ws.delete_rows(r, 1)

    # ---- 3. comments do not survive a row delete cleanly: re-apply ours
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            if ws.cell(row, col).comment:
                ws.cell(row, col).comment = None
    index = {(s, d): r for r, s, d in district_rows(ws)}
    for (state, child), parent in MERGE.items():
        pr = index[(state, parent)]
        note = (f"{child} is a court establishment inside {parent}, not a "
                f"district of its own. Its row has been merged into this one: "
                f"the judge counts already covered {child} (both were read "
                f"from the same court website), and its recorded counts "
                f"have been added here.")
        for col in (2, 3, 4):
            ws.cell(pr, col).comment = Comment(note, "judge count")
            ws.cell(pr, col).comment.width = 340
            ws.cell(pr, col).comment.height = 150
    for key, note in NOTES.items():
        r = index.get(key)
        if r:
            for col in (2, 3, 4):
                ws.cell(r, col).comment = Comment(note, "judge count")
                ws.cell(r, col).comment.width = 340
                ws.cell(r, col).comment.height = 150

    # ---- 4. rebuild the formulas the deletes invalidated
    for r, _s, _d in district_rows(ws):
        ws.cell(r, 12).value = f'=IFERROR(F{r}/E{r},"")'
        ws.cell(r, 13).value = f'=IFERROR(G{r}/F{r},"")'
        ws.cell(r, 14).value = f'=IFERROR(I{r}/G{r},"")'
    blocks = state_blocks(ws)
    for _state, lo, hi, tr in blocks:
        for c in range(3, 12):
            col = chr(64 + c)
            ws.cell(tr, c).value = f"=SUM({col}{lo}:{col}{hi})"
        ws.cell(tr, 12).value = f'=IFERROR(F{tr}/E{tr},"")'
        ws.cell(tr, 13).value = f'=IFERROR(G{tr}/F{tr},"")'
        ws.cell(tr, 14).value = f'=IFERROR(I{tr}/G{tr},"")'

    # ---- 5. recompute the States sheet from the Districts sheet
    st = wb["States"]
    strow = {str(st.cell(r, 1).value): r for r in range(2, st.max_row + 1)
             if st.cell(r, 1).value}
    state_changes = []
    for state, lo, hi, _tr in blocks:
        sr = strow.get(state)
        if not sr:
            continue
        with_accounts, sums = 0, [0] * len(COUNT_COLS)
        for r in range(lo, hi + 1):
            d = ws.cell(r, 2).value
            acc = ws.cell(r, 5).value
            for i, c in enumerate(COUNT_COLS):
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)):
                    sums[i] += v
            if d and str(d) != "Not recorded" and isinstance(acc, (int, float)) and acc > 0:
                with_accounts += 1
        old = [st.cell(sr, 2).value] + [st.cell(sr, c).value for c in range(6, 13)]
        st.cell(sr, 2).value = with_accounts
        for i, c in enumerate(range(6, 13)):
            st.cell(sr, c).value = sums[i]
        new = [with_accounts] + sums
        if old != new:
            state_changes.append((state, old, new))

    wb.save(args.out)

    # ---- report
    print(f"wrote {args.out}\n")
    print("merged into parent district (judges already counted there):")
    for state, child, parent, moved, pc, pd in merged:
        print(f"  {state} / {child} -> {parent}   parent C={pc} D={pd}")
        print(f"      accounts carried across: {moved}")
    print("\nremoved outright (filed under the wrong state):")
    for key, why in DROP.items():
        f = dropped_counts.get(key)
        print(f"  {key[0]} / {key[1]}: {why}")
        print(f"      count values dropped E..K: {f}")
    print(f"\nrows on Districts sheet: {ws.max_row} (7 fewer)")
    print("\nStates sheet rows updated:")
    hdr = ["Districts w/ accounts"] + [st.cell(1, c).value for c in range(6, 13)]
    for state, old, new in state_changes:
        diffs = [f"{h}: {o} -> {n}" for h, o, n in zip(hdr, old, new) if o != n]
        print(f"  {state}: " + "; ".join(diffs))


if __name__ == "__main__":
    main()
