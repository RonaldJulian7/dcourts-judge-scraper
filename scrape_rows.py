#!/usr/bin/env python3
"""
Stage 1: pull every officer row from the district court sites (*.dcourts.gov.in).

Writes one CSV row per officer -- no counting, no classification. Stage 2
(classify.py) turns this into the per-district totals, so the judgement calls
about what counts as "bench" can be changed without re-scraping.

    python scrape_rows.py --xlsx districts.xlsx

Why it is built this way: the /list-of-judges/ pages are WordPress templates
that render one <div class="whoswho"> block per court complex, each with its
own table. Blocks longer than 20 rows are truncated in the static HTML and the
rest is fetched by the theme's own AJAX call (action=judge_pages), so this
follows that call rather than parsing only what the first response contains.

Outputs:
    judges_rows.csv     state, district, url, block, name, designation
    judges_sites.csv    one row per district: which URL worked, block/row
                        counts, whether any AJAX page was missed, status
"""

import argparse
import csv
import hashlib
import json
import os
import re
import sys
import threading
import time
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor, as_completed

try:
    import requests
    from bs4 import BeautifulSoup
except ImportError:
    sys.exit("pip install requests beautifulsoup4 lxml openpyxl")

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/122.0 Safari/537.36")

PATHS = ["/list-of-judges/", "/judges/", "/hon-ble-judges/",
         "/list-of-judicial-officers/", "/judicial-officers/"]

# district name -> subdomain slug, where the obvious guess is wrong.
# Everything below the first group was found by DNS-probing candidates after
# the first full run reported them as not_found.
OVERRIDES = {
    ("Andhra Pradesh", "West Godavari"): "westgodavari",
    ("Andhra Pradesh", "East Godavari"): "eastgodavari",
    ("Punjab", "S.A.S. Nagar"): "sasnagar",
    ("Punjab", "S.B.S. Nagar"): "sbsnagar",
    ("Punjab", "Sri Muktsar Sahib"): "muktsarsahib",
    ("Haryana", "Yamuna Nagar"): "yamunanagar",
    # the court sits at the district headquarters town, not the district name
    ("Haryana", "Mahendragarh"): "narnaul",

    ("Karnataka", "Dakshina Kannada"): "dk",
    ("Karnataka", "Chamarajanagar"): "chamarajanagara",
    ("Karnataka", "Vijayanagara"): "vijayanagar",
    ("Madhya Pradesh", "Narmadapuram"): "narmadapuram",
    ("Odisha", "Kendrapada"): "kendrapara",
    ("Odisha", "Anugul"): "angul",
    ("Andhra Pradesh", "Visakapatnam"): "visakhapatnam",
    ("Andhra Pradesh", "Ananthapur"): "anantapur",
    ("Andhra Pradesh", "Prakasham"): "prakasam",
    ("Andhra Pradesh", "Nellore"): "spsrnellore",
    ("Haryana", "Sonipat"): "sonepat",
    ("Rajasthan", "Jaipur"): "jaipurdistrict",
    ("Rajasthan", "Jaipur Metro I"): "jaipurmetro1",
    ("Rajasthan", "Jaipur Metro II"): "jaipurmetro2",
    ("Rajasthan", "Merta Nagaur"): "nagaur",
    ("Rajasthan", "Salumber"): "salumbar",
    ("Delhi", "Central"): "centraldelhi",
    ("Delhi", "West"): "westdelhi",
    ("Delhi", "North"): "northdelhi",
    ("Delhi", "South"): "southdelhi",
    ("Delhi", "South East"): "southeastdelhi",
    ("Bihar", "Kaimur at Bhabhua"): "bhabua",
    # aurangabad.dcourts.gov.in is MAHARASHTRA's Aurangabad. Bihar's sits on a
    # state-qualified host. Found by checking every site's home-page <title>
    # against the state the workbook expects.
    ("Bihar", "Aurangabad"): "aurangabad.bih",
}

# Rows in the workbook that are a court establishment inside another
# district's site, not a judicial district of their own. Take only the blocks
# on the parent site whose name matches, so the parent's officers are not
# counted twice.
SUBSITES = {
    ("Andhra Pradesh", "Bheemunipatnam"): ("visakhapatnam", r"bheemuni|bhimuni"),
    ("Andhra Pradesh", "Parchur"): ("prakasam", r"parchur"),
    ("Madhya Pradesh", "Berasia"): ("bhopal", r"berasia"),
}

HDR_PAT = re.compile(
    r"^(sr|s\.?\s*no|sl|serial|no\.?$|name|designation|court|from|to|period|"
    r"posting|remarks|photo|profile|sitting|present|room|location)\b", re.I)
NOISE_PAT = re.compile(r"^\s*(nil|n\.?a\.?|-+|â€”+|not\s+posted)\s*$", re.I)
# A vacant post is not an officer posted. The sites write this several ways,
# including "Post Kept Vacant ( Incharge ) <designation of the covering
# officer>" -- counting those would also double-count the covering officer.
VACANT_PAT = re.compile(r"\bvacant\b|\bvacancy\b|post\s+kept", re.I)

_print_lock = threading.Lock()


def log(msg):
    with _print_lock:
        print(msg, flush=True)


# --------------------------------------------------------------- fetching
def slugs_for(state, district):
    if (state, district) in SUBSITES:
        return [SUBSITES[(state, district)][0]]
    if (state, district) in OVERRIDES:
        return [OVERRIDES[(state, district)]]
    base = re.sub(r"\(.*?\)", " ", district.lower())
    base = re.sub(r"[^a-z0-9 ]", " ", base)
    words = base.split()
    out = ["".join(words)]
    if len(words) > 1:
        out += ["-".join(words), words[0]]
    seen, uniq = set(), []
    for s in out:
        if s and s not in seen:
            seen.add(s)
            uniq.append(s)
    return uniq


def cache_path(cache_dir, key):
    return os.path.join(cache_dir, hashlib.sha1(key.encode()).hexdigest()[:16] + ".txt")


class Fetcher:
    """GET/POST with an on-disk cache. Returns (status, text, from_cache)."""

    def __init__(self, cache_dir, refresh, retries, delay, timeout):
        self.cache_dir = cache_dir
        self.refresh = refresh
        self.retries = retries
        self.delay = delay
        self.timeout = timeout
        self.local = threading.local()

    @property
    def session(self):
        s = getattr(self.local, "s", None)
        if s is None:
            s = requests.Session()
            s.headers.update({"User-Agent": UA, "Accept-Language": "en-IN,en;q=0.9"})
            self.local.s = s
        return s

    def _cached(self, key):
        cp = cache_path(self.cache_dir, key)
        if not self.refresh and os.path.exists(cp):
            with open(cp, encoding="utf-8", errors="replace") as fh:
                return fh.read()
        return None

    def _store(self, key, text):
        with open(cache_path(self.cache_dir, key), "w", encoding="utf-8") as fh:
            fh.write(text)

    def get(self, url):
        hit = self._cached(url)
        if hit is not None:
            return 200, hit, True
        last = None
        for attempt in range(self.retries + 1):
            try:
                r = self.session.get(url, timeout=self.timeout, allow_redirects=True)
                if r.status_code == 200 and r.text:
                    if not r.encoding or r.encoding.lower() == "iso-8859-1":
                        r.encoding = r.apparent_encoding or "utf-8"
                    self._store(url, r.text)
                    return 200, r.text, False
                last = r.status_code
                if last in (404, 403, 410):     # settled answer, do not retry
                    return last, None, False
            except requests.exceptions.SSLError:
                return "ssl", None, False
            except requests.exceptions.ConnectionError as exc:
                # DNS failure means the subdomain does not exist at all -- no
                # point trying the other paths on this host.
                if "NameResolutionError" in str(exc) or "getaddrinfo" in str(exc):
                    return "dns", None, False
                last = "conn"
            except requests.RequestException:
                last = "timeout"
            if attempt < self.retries:
                time.sleep(self.delay * (attempt + 1))
        return last, None, False

    def post(self, url, data):
        key = url + "|" + json.dumps(data, sort_keys=True)
        hit = self._cached(key)
        if hit is not None:
            return 200, hit, True
        last = None
        for attempt in range(self.retries + 1):
            try:
                r = self.session.post(url, data=data, timeout=self.timeout,
                                      headers={"X-Requested-With": "XMLHttpRequest"})
                if r.status_code == 200 and r.text:
                    self._store(key, r.text)
                    return 200, r.text, False
                last = r.status_code
            except requests.RequestException:
                last = "timeout"
            if attempt < self.retries:
                time.sleep(self.delay * (attempt + 1))
        return last, None, False


# --------------------------------------------------------------- parsing
def strip_html(val):
    if "<" in val:
        val = BeautifulSoup(val, "lxml").get_text(" ", strip=True)
    return re.sub(r"\s+", " ", val).strip()


def header_map(table):
    """column index -> header label, from <thead> or the first <th> row."""
    head = table.find("thead")
    tr = head.find("tr") if head else None
    if tr is None:
        for cand in table.find_all("tr"):
            if cand.find("th"):
                tr = cand
                break
    if tr is None:
        return {}
    return {i: re.sub(r"\s+", " ", c.get_text(" ", strip=True))
            for i, c in enumerate(tr.find_all(["th", "td"]))}


def pick(cells, hdr):
    """(name, designation) using the header labels, falling back to position."""
    name = desig = None
    extra = []
    for i, val in enumerate(cells):
        label = hdr.get(i, "").lower()
        if not val:
            continue
        if "profile" in label or "photo" in label or "picture" in label:
            continue
        if re.fullmatch(r"(sr|sl)?\.?\s*no\.?", label.strip()):
            continue
        if "name" in label and name is None:
            name = val
        elif "designation" in label and desig is None:
            desig = val
        elif label:
            extra.append(val)
        else:
            extra.append(val)
    if name is None or desig is None:
        # no usable header: drop a leading serial cell, then take first two
        vals = [v for v in cells if v]
        if vals and re.fullmatch(r"\d{1,3}\.?", vals[0]):
            vals = vals[1:]
        vals = [v for v in vals if not re.fullmatch(r"\d{1,4}", v)]
        if len(vals) < 2:
            return None, None, []
        name = name or vals[0]
        desig = desig or " | ".join(vals[1:])
        extra = []
    return name, desig, extra


def keep_row(name, desig):
    """the one place that decides whether a parsed row is a posted officer."""
    if not name or not desig:
        return False
    if len(name) < 3 or NOISE_PAT.match(name) or VACANT_PAT.search(name):
        return False
    if HDR_PAT.match(name) and len(name) < 30:
        return False
    return True


def rows_from_table(table):
    hdr = header_map(table)
    body = table.find("tbody") or table
    out = []
    for tr in body.find_all("tr"):
        if tr.find_parent("thead"):
            continue
        tds = tr.find_all(["td", "th"])
        if tr.find("th") and not tr.find("td"):
            continue                              # header row
        cells = [re.sub(r"\s+", " ", c.get_text(" ", strip=True)) for c in tds]
        if len([c for c in cells if c]) < 2:
            continue
        name, desig, _ = pick(cells, hdr)
        if keep_row(name, desig):
            out.append((name, desig))
    return out


def ajax_rows(text):
    """
    -> (rows, is_last, paged, n_raw). n_raw is the number of records the server
    sent before filtering, so the caller can tell "server has no more rows"
    from "every row on this page was a vacancy".
    """
    try:
        obj = json.loads(text)
    except ValueError:
        return [], True, None, 0
    items = obj.get("html") or []
    if isinstance(items, dict):
        items = list(items.values())
    out = []
    for it in items:
        if not isinstance(it, dict):
            continue
        name = desig = None
        for k, v in it.items():
            kl = str(k).lower()
            v = strip_html(str(v))
            if "profile" in kl or "picture" in kl or "photo" in kl:
                continue
            if "name" in kl and name is None:
                name = v
            elif "designation" in kl and desig is None:
                desig = v
        if keep_row(name, desig):
            out.append((name, desig))
    return out, bool(obj.get("last")), obj.get("paged"), len(items)


def parse_district(html, url, fetcher, delay):
    """
    -> (rows, blocks, note) where rows are (block_name, name, designation).

    Long lists are truncated at 20 rows and continued two different ways
    depending on whether the block belongs to a taxonomy:
      - blocks with a data-id continue through the theme's judge_pages AJAX
      - uncategorised blocks continue at ordinary /list-of-judges/page/N/ URLs
    Both are followed here. Rows are de-duplicated, so a block that repeats on
    every page is counted once.
    """
    rows, blocks, notes, seen = [], [], [], set()

    def add(block, name, desig):
        key = (block, name.lower(), desig.lower())
        if key not in seen:
            seen.add(key)
            rows.append((block, name, desig))

    pending, visited, guard = [(url, html)], {url}, 0
    while pending and guard < 30:
        guard += 1
        u, h = pending.pop(0)
        for nu in parse_one(h, u, fetcher, delay, add, blocks, notes):
            if nu in visited:
                continue
            visited.add(nu)
            st, nh, cached = fetcher.get(nu)
            if not cached:
                time.sleep(delay)
            if nh:
                pending.append((nu, nh))
            else:
                notes.append(f"page fetch failed ({st}): {nu}")

    uniq = OrderedDict()
    for bname, ns, na, pag in blocks:
        cur = uniq.get(bname, (0, 0, False))
        uniq[bname] = (cur[0] + ns, cur[1] + na, cur[2] or pag)
    blocks = [(b, v[0], v[1], v[2]) for b, v in uniq.items()]
    return rows, blocks, "; ".join(notes)


def parse_one(html, url, fetcher, delay, add, blocks, notes):
    """parse a single page; returns further page URLs to follow."""
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "nav", "header", "footer"]):
        tag.decompose()

    origin = re.match(r"(https?://[^/]+)", url).group(1)
    ajaxurl = origin + "/wp-admin/admin-ajax.php"
    m = re.search(r'"ajaxUrl"\s*:\s*"([^"]+admin-ajax\.php)"', html)
    if m:
        ajaxurl = m.group(1).replace("\\/", "/")

    wwh = soup.select_one(".ww_html")
    nonce = wwh.get("data-nonce") if wwh else None
    profile = (wwh.get("data-profile") if wwh else None) or "1"
    dfilter = (wwh.get("data-filter") if wwh else None) or "name,designation"
    sl = soup.select_one("#show_link_or_not")
    show_link = (sl.get("value") if sl else None) or "0"

    # real hrefs in the pagination mean server-side paging, which is what the
    # uncategorised blocks use
    next_urls = []
    for a in soup.select(".pagination a[href]"):
        href = a["href"]
        if "/page/" in href and not href.startswith("#"):
            next_urls.append(requests.compat.urljoin(url, href))

    ww = soup.select("div.whoswho")
    if ww:
        for b in ww:
            table = b.find("table")
            cap = table.find("caption") if table else None
            bname = ((b.get("data-name") or "").strip()
                     or (cap.get_text(" ", strip=True) if cap else "")
                     or "uncategorized")
            bid = b.get("data-id")
            got = rows_from_table(table) if table else []
            for n, d in got:
                add(bname, n, d)
            n_static = len(got)
            n_ajax = 0
            pag = b.select_one(".pagination")
            if pag and bid and nonce:
                paged, guard = 1, 0
                while guard < 25:
                    guard += 1
                    st, txt, cached = fetcher.post(ajaxurl, {
                        "action": "judge_pages", "paged": str(paged),
                        "tax_slug": str(bid), "type": "next",
                        "profile_image": str(profile), "filter": str(dfilter),
                        "show_link": str(show_link), "security": str(nonce)})
                    if not cached:
                        time.sleep(delay)
                    if txt is None:
                        notes.append(f"ajax failed for block '{bname}' ({st})")
                        break
                    got2, is_last, newpaged, n_raw = ajax_rows(txt)
                    if not n_raw:
                        break
                    for n, d in got2:
                        add(bname, n, d)
                    n_ajax += len(got2)
                    if is_last:
                        break
                    paged = newpaged or (paged + 1)
            elif pag and not nonce:
                notes.append(f"block '{bname}' is paginated but no nonce found")
            blocks.append((bname, n_static, n_ajax, bool(pag)))
    else:
        # not the whoswho template: take every table on the page
        found = 0
        for i, table in enumerate(soup.find_all("table")):
            cap = table.find("caption")
            bname = (cap.get_text(" ", strip=True) if cap else f"table{i + 1}")
            got = rows_from_table(table)
            for n, d in got:
                add(bname, n, d)
            found += len(got)
            blocks.append((bname, len(got), 0, False))
        if not found:
            notes.append("pdf link on page" if re.search(r'\.pdf"', html, re.I)
                         else "no table found")
    return next_urls


# --------------------------------------------------------------- districts
def districts_from_xlsx(path, sheet="Districts", extra_sheet=""):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    out = OrderedDict()
    for row in ws.iter_rows(min_row=2, values_only=True):
        state, district = row[0], row[1]
        if not state or not district:
            continue
        d = str(district)
        if d.startswith("Total") or d == "Not recorded":
            continue
        # A row whose count column holds an em dash is not a judicial district
        # at all -- an entry filed under a district name that does not exist in
        # that state. Leave them alone.
        posted = row[2] if len(row) > 2 else None
        if isinstance(posted, str) and posted.strip() in ("â€”", "â€“", "-"):
            continue
        out.setdefault(str(state), [])
        if d not in out[str(state)]:
            out[str(state)].append(d)

    # An optional second tab may list the same two columns for a wider set of
    # districts than the main one covers. Merge its districts in so every
    # sheet can be filled from a single pass. It is expected to hold the state
    # in column B and the district in column C.
    if extra_sheet and extra_sheet in wb.sheetnames:
        ph = wb[extra_sheet]
        for row in ph.iter_rows(min_row=2, values_only=True):
            if len(row) < 3:
                continue
            state, district = row[1], row[2]
            if not state or not district:
                continue
            d = str(district)
            out.setdefault(str(state), [])
            if d not in out[str(state)]:
                out[str(state)].append(d)
    return out


# --------------------------------------------------------------- one district
def do_district(state, district, fetcher, delay):
    site = {"state": state, "district": district, "url": "", "blocks": 0,
            "rows": 0, "paginated_blocks": 0, "status": "not_found", "note": "",
            "tried": ""}
    tried, codes = [], []
    sub = SUBSITES.get((state, district))
    block_filter = re.compile(sub[1], re.I) if sub else None
    for slug in slugs_for(state, district):
        dead = False
        for path in PATHS:
            url = f"https://{slug}.dcourts.gov.in{path}"
            tried.append(url)
            st, html, cached = fetcher.get(url)
            if not cached:
                time.sleep(delay)
            if html is None:
                codes.append(str(st))
                if st in ("dns", "ssl"):
                    dead = True          # whole host is gone, skip other paths
                    break
                continue
            rows, blocks, note = parse_district(html, url, fetcher, delay)
            if block_filter is not None:
                # some states group blocks by cadre rather than by court
                # location, so the establishment shows up in the designation
                kept = [r for r in rows if block_filter.search(r[0] + " " + r[2])]
                note = (note + "; " if note else "") + (
                    f"establishment inside {slug}: kept {len(kept)} of {len(rows)} rows"
                    if kept else
                    f"no block matching '{sub[1]}' on {slug} site")
                rows = kept
            site["url"] = url
            site["blocks"] = len(blocks)
            site["rows"] = len(rows)
            site["paginated_blocks"] = sum(1 for b in blocks if b[3])
            site["note"] = note
            if rows:
                site["status"] = "ok"
                return site, [(state, district, url, b, n, d) for b, n, d in rows]
            site["status"] = "no_rows"
        if dead:
            continue
    site["status"] = site["status"] if site["url"] else "not_found"
    site["tried"] = tried[0] if tried else ""
    if not site["note"]:
        site["note"] = "codes: " + ",".join(sorted(set(codes))[:5])
    return site, []


# --------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True,
                    help="workbook listing the districts to scrape")
    ap.add_argument("--sheet", default="Districts",
                    help="sheet holding state in column A, district in column B")
    ap.add_argument("--extra-sheet", default="",
                    help="optional second sheet with state in B, district in C")
    ap.add_argument("--rows-out", default="judges_rows.csv")
    ap.add_argument("--sites-out", default="judges_sites.csv")
    ap.add_argument("--cache", default=".dcourts_cache2")
    ap.add_argument("--delay", type=float, default=0.4)
    ap.add_argument("--retries", type=int, default=2)
    ap.add_argument("--timeout", type=float, default=45)
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--states", default="")
    ap.add_argument("--districts", default="")
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--refresh", action="store_true")
    args = ap.parse_args()

    os.makedirs(args.cache, exist_ok=True)
    tree = districts_from_xlsx(args.xlsx, args.sheet, args.extra_sheet)
    if args.states:
        keep = {s.strip() for s in args.states.split(",")}
        tree = OrderedDict((k, v) for k, v in tree.items() if k in keep)
    targets = [(s, d) for s, ds in tree.items() for d in ds]
    if args.districts:
        want = {tuple(x.strip() for x in it.split(":", 1))
                for it in args.districts.split(",") if ":" in it}
        targets = [t for t in targets if t in want]
    if args.limit:
        targets = targets[:args.limit]

    fetcher = Fetcher(args.cache, args.refresh, args.retries, args.delay, args.timeout)
    sites, all_rows, done = [], [], [0]
    total = len(targets)

    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futs = {pool.submit(do_district, s, d, fetcher, args.delay): (s, d)
                for s, d in targets}
        for fut in as_completed(futs):
            s, d = futs[fut]
            try:
                site, rows = fut.result()
            except Exception as exc:
                site = {"state": s, "district": d, "url": "", "blocks": 0,
                        "rows": 0, "paginated_blocks": 0, "status": "error",
                        "note": f"{type(exc).__name__}: {exc}"[:200], "tried": ""}
                rows = []
            sites.append(site)
            all_rows.extend(rows)
            done[0] += 1
            log(f"[{done[0]}/{total}] {s} / {d}: {site['status']} "
                f"rows={site['rows']} blocks={site['blocks']}"
                + (f"  {site['note']}" if site['note'] else ""))

    order = {t: i for i, t in enumerate(targets)}
    sites.sort(key=lambda r: order.get((r["state"], r["district"]), 1 << 30))
    all_rows.sort(key=lambda r: order.get((r[0], r[1]), 1 << 30))

    with open(args.rows_out, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["state", "district", "url", "block", "name", "designation"])
        w.writerows(all_rows)

    cols = ["state", "district", "url", "blocks", "rows", "paginated_blocks",
            "status", "note", "tried"]
    with open(args.sites_out, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        w.writerows(sites)

    ok = [s for s in sites if s["status"] == "ok"]
    log(f"\n{len(ok)}/{len(sites)} districts scraped, {len(all_rows)} officer rows")
    log(f"wrote {args.rows_out} and {args.sites_out}")
    bad = [s for s in sites if s["status"] != "ok"]
    if bad:
        log(f"{len(bad)} districts need a manual look:")
        for s in bad:
            log(f"   {s['state']} / {s['district']}: {s['status']} {s['note']}")


if __name__ == "__main__":
    main()
