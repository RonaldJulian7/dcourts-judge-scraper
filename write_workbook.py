#!/usr/bin/env python3
"""
Stage 3: write columns C and D of the Districts sheet into a copy of the
workbook, and report every value that changed.

The source workbook is never modified -- output goes to --out. Formulas
elsewhere on the sheet (the L/M/N ratios, the SUM rows) are loaded as formulas
and written back untouched; the "Total <State>" row for column D gets a SUM
mirroring the range its column C formula already uses.

Districts with no scraped data keep whatever the sheet already had, and are
listed at the end so nothing changes silently.

    python write_workbook.py --xlsx districts.xlsx \
        --summary judges_summary.csv --out districts_filled.xlsx
"""

import argparse
import csv
import re
import shutil

# Revenue districts with no judicial district of their own: nobody is posted
# to them, so both columns are 0 and the reason goes into a cell note rather
# than being left for the reader to guess.
#
# Evidence for each: Sikkim's High Court lists exactly four district courts
# (Mangan, Gangtok, Namchi, Gyalshing); no pakyong/soreng host exists on
# dcourts.gov.in; neither name appears in any of the scraped judge lists; and
# the cause-list court-establishment dropdowns on the Gangtok and Gyalshing
# sites show every establishment sitting at the district headquarters, with
# none at Pakyong or Soreng.
ZERO_WITH_NOTE = {
    ("Sikkim", "Pakyong"): (
        "Not a judicial district. Pakyong was carved out of East Sikkim in "
        "2021 but has no court of its own: it falls under the Gangtok "
        "judicial district (District & Sessions Court, Gangtok). All four "
        "Gangtok establishments - District & Sessions Judge, Civil Judge Sr. "
        "Division, Chief Judicial Magistrate and Family Court - sit at "
        "Gangtok, so the officers serving Pakyong are counted in the Gangtok "
        "row above."),
    ("Sikkim", "Soreng"): (
        "Not a judicial district. Soreng was carved out of West Sikkim in "
        "2021 but has no court of its own: it falls under the Gyalshing "
        "judicial district (District & Sessions Court, Gyalshing). All four "
        "Gyalshing establishments - District & Sessions Judge, Civil Judge "
        "Sr. Division, Chief Judicial Magistrate and Family Court - sit at "
        "Gyalshing, so the officers serving Soreng are counted in the "
        "Gyalshing row below."),
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--summary", default="judges_summary.csv")
    ap.add_argument("--out", default="districts_filled.xlsx")
    ap.add_argument("--changes-out", default="judges_changes.csv")
    ap.add_argument("--sheet", default="Districts",
                    help="main sheet: state in A, district in B, counts in C/D")
    ap.add_argument("--extra-sheet", default="",
                    help="optional second sheet: state in B, district in C, "
                         "counts in D/E")
    args = ap.parse_args()

    from openpyxl import load_workbook
    from openpyxl.comments import Comment

    summary = {(r["state"], r["district"]): r
               for r in csv.DictReader(open(args.summary, encoding="utf-8-sig"))}

    shutil.copyfile(args.xlsx, args.out)
    wb = load_workbook(args.out, data_only=False)
    ws = wb[args.sheet]

    changes, skipped, marked, totals = [], [], [], 0
    for row in range(2, ws.max_row + 1):
        state = ws.cell(row, 1).value
        district = ws.cell(row, 2).value
        if not district:
            continue
        d = str(district)

        if d.startswith("Total"):
            # mirror column C's SUM range onto column D
            cf = ws.cell(row, 3).value
            m = re.match(r"=SUM\(C(\d+):C(\d+)\)", str(cf) or "")
            if m:
                ws.cell(row, 4).value = f"=SUM(D{m.group(1)}:D{m.group(2)})"
                ws.cell(row, 4).number_format = ws.cell(row, 3).number_format
                totals += 1
            continue

        if not state or d == "Not recorded":
            continue

        old_c, old_d = ws.cell(row, 3).value, ws.cell(row, 4).value
        # an em dash in column C marks a row that is not a judicial district
        if isinstance(old_c, str) and old_c.strip() in ("—", "–", "-"):
            continue

        if (str(state), d) in ZERO_WITH_NOTE:
            note = ZERO_WITH_NOTE[(str(state), d)]
            ws.cell(row, 3).value = 0
            ws.cell(row, 4).value = 0
            ws.cell(row, 4).number_format = ws.cell(row, 3).number_format
            for col in (2, 3, 4):
                c = ws.cell(row, col)
                c.comment = Comment(note, "judge count")
                c.comment.width = 340
                c.comment.height = 150
            marked.append((str(state), d, old_c, old_d, note))
            continue

        rec = summary.get((str(state), d))
        if not rec or rec["officers_posted"] == "":
            skipped.append((str(state), d, old_c, old_d,
                            rec["status"] if rec else "not in summary"))
            continue

        new_c = int(rec["officers_posted"])
        new_d = int(rec["bench"])
        ws.cell(row, 3).value = new_c
        ws.cell(row, 4).value = new_d
        ws.cell(row, 4).number_format = ws.cell(row, 3).number_format
        changes.append({"state": state, "district": d, "row": row,
                        "old_posted": old_c, "new_posted": new_c,
                        "old_bench": old_d, "new_bench": new_d,
                        "diff_posted": (new_c - old_c)
                        if isinstance(old_c, int) else "",
                        "non_bench": rec["non_bench"],
                        "distinct_persons": rec["distinct_persons"],
                        "url": rec["url"]})

    # ---- optional second sheet: same two columns, one column further right ----
    ph_changes, ph_skipped = [], []
    if args.extra_sheet and args.extra_sheet in wb.sheetnames:
        ph = wb[args.extra_sheet]
        for row in range(2, ph.max_row + 1):
            state, district = ph.cell(row, 2).value, ph.cell(row, 3).value
            if not state or not district:
                continue                       # group headings and total rows
            rec = summary.get((str(state), str(district)))
            old_c, old_d = ph.cell(row, 4).value, ph.cell(row, 5).value
            if not rec or rec["officers_posted"] == "":
                ph_skipped.append((str(state), str(district), old_c, old_d))
                continue
            ph.cell(row, 4).value = int(rec["officers_posted"])
            ph.cell(row, 5).value = int(rec["bench"])
            ph.cell(row, 5).number_format = ph.cell(row, 4).number_format
            ph_changes.append({"state": state, "district": district, "row": row,
                               "old_posted": old_c,
                               "new_posted": int(rec["officers_posted"]),
                               "old_bench": old_d, "new_bench": int(rec["bench"]),
                               "diff_posted": (int(rec["officers_posted"]) - old_c)
                               if isinstance(old_c, int) else "",
                               "non_bench": rec["non_bench"],
                               "distinct_persons": rec["distinct_persons"],
                               "url": rec["url"]})

    wb.save(args.out)

    cols = ["state", "district", "row", "old_posted", "new_posted",
            "diff_posted", "old_bench", "new_bench", "non_bench",
            "distinct_persons", "url"]
    with open(args.changes_out, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=["sheet"] + cols)
        w.writeheader()
        w.writerows([{"sheet": args.sheet, **c} for c in changes])
        w.writerows([{"sheet": args.extra_sheet, **c} for c in ph_changes])

    print(f"wrote {args.out}")
    print(f"{args.sheet} sheet")
    print(f"  districts filled       : {len(changes)}")
    print(f"  Total rows given a SUM : {totals}")
    same = sum(1 for c in changes if c["diff_posted"] == 0)
    print(f"  column C unchanged     : {same}")
    print(f"  column C changed       : {len(changes) - same}")
    if marked:
        print(f"  set to 0 with a cell note: {len(marked)}")
        for s in marked:
            print(f"      {s[0]} / {s[1]}: C={s[2]!r} D={s[3]!r} -> 0 / 0")
            print(f"          note: {s[4][:88]}...")
    if skipped:
        print(f"  left alone (no data)   : {len(skipped)}")
        for s in skipped:
            print(f"      {s[0]} / {s[1]}: kept C={s[2]!r} D={s[3]!r}  ({s[4]})")
    print(f"{args.extra_sheet or 'second'} sheet")
    print(f"  districts filled       : {len(ph_changes)}")
    same_ph = sum(1 for c in ph_changes if c["diff_posted"] == 0)
    print(f"  column D unchanged     : {same_ph}")
    print(f"  column D changed       : {len(ph_changes) - same_ph}")
    if ph_skipped:
        print(f"  left alone (no data)   : {len(ph_skipped)}")
        for s in ph_skipped:
            print(f"      {s[0]} / {s[1]}: kept D={s[2]!r} E={s[3]!r}")
    print(f"per-district detail      : {args.changes_out}")


if __name__ == "__main__":
    main()
